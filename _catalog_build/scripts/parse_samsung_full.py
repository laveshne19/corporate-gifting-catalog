# -*- coding: utf-8 -*-
import openpyxl, os, json

BASE = "/Users/laveshbansal/Downloads/📁 Master Folder/master price list"
OUT = os.path.join(BASE, "_catalog_build", "extracted")
F = "Samsung Specs, SKU, SNS, SC+ and KNOX (2).xlsx"

def num(v):
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return round(float(v), 2)
    return None

wb = openpyxl.load_workbook(os.path.join(BASE, F), read_only=True, data_only=True)
records = []

def clean(v):
    return str(v).strip() if v is not None else ""

# ============ Tablets / Smartphones (same layout) ============
def parse_device_sheet(sheet_name, category):
    ws = wb[sheet_name]
    out = []
    cur_subcat = ""
    for row in list(ws.iter_rows(values_only=True))[3:]:
        subcat_cell, name, code, status, eop, upg, trans, mrp, mop = (list(row) + [None]*9)[1:10]
        if subcat_cell:
            cur_subcat = clean(subcat_cell)
        if not name or not code:
            continue
        specs_bits = []
        row_l = list(row)
        # chipset(10) ram(13) rom(14) inch(16) camera_rear(20)
        for idx, label in [(10, "Chipset"), (13, "RAM"), (14, "Storage"), (16, "Display"), (20, "Camera")]:
            if idx < len(row_l) and row_l[idx]:
                specs_bits.append(f"{label}: {clean(row_l[idx])}")
        out.append(dict(
            brand="Samsung", category=category, sub_category=cur_subcat,
            product_name=f"Galaxy {clean(name)}", model_code=clean(code),
            description=f"Galaxy {clean(name)} | " + " | ".join(specs_bits),
            variant="", key_specs=" | ".join(specs_bits),
            mrp=num(mrp), landing_price=num(mop), warranty="",
            source_file=F, source_sheet=sheet_name,
        ))
    return out

records += parse_device_sheet("Tablets", "Mobiles & Tablets")
records += parse_device_sheet("Smartphones", "Mobiles & Tablets")

# ============ Wearables ============
def parse_wearables():
    ws = wb["Wearables"]
    out = []
    cur_subcat = ""
    for row in list(ws.iter_rows(values_only=True))[3:]:
        subcat_cell, name, code, status, eop, upg, trans, mrp, mop = (list(row) + [None]*9)[1:10]
        if subcat_cell:
            cur_subcat = clean(subcat_cell)
        if not name or not code:
            continue
        out.append(dict(
            brand="Samsung", category="Wearables", sub_category=cur_subcat,
            product_name=f"Galaxy {clean(name)}", model_code=clean(code),
            description=f"Galaxy {clean(name)}", variant="", key_specs=cur_subcat,
            mrp=num(mrp), landing_price=num(mop), warranty="",
            source_file=F, source_sheet="Wearables",
        ))
    return out

records += parse_wearables()

# ============ Accessories ============
def parse_acc():
    ws = wb["Acc Pricing"]
    out = []
    for row in list(ws.iter_rows(values_only=True))[1:]:
        sno, acc_cat, device, item, model, sku, specs, mrp, mop = (list(row) + [None]*9)[:9]
        if not item or mrp is None:
            continue
        name = f"{clean(item)} ({clean(specs)})" if specs else clean(item)
        out.append(dict(
            brand="Samsung", category="Mobile Accessories", sub_category=clean(acc_cat),
            product_name=name, model_code=clean(model), description=f"{name} — for {clean(device)}",
            variant=clean(specs), key_specs=f"For {clean(device)}",
            mrp=num(mrp), landing_price=num(mop), warranty="",
            source_file=F, source_sheet="Acc Pricing",
        ))
    return out

records += parse_acc()

# ============ NPC (laptops) — supersedes earlier partial Galaxy Book data ============
def parse_npc():
    ws = wb["NPC Specs"]
    out = []
    for row in list(ws.iter_rows(values_only=True))[1:]:
        name, sku, india_sku, color, screen, cpu, ram, storage, os_, gfx, mrp, mop, remarks, model_type = (list(row) + [None]*14)[:14]
        if not name or not sku or mrp is None:
            continue
        specs = f"{clean(cpu)}, {clean(ram)} RAM, {clean(storage)} Storage, {clean(screen)}, {clean(os_)}"
        out.append(dict(
            brand="Samsung", category="Laptops & Computing", sub_category=clean(name),
            product_name=f"{clean(name)} ({clean(sku)})", model_code=clean(sku),
            description=f"{clean(name)} — {specs}", variant=clean(color), key_specs=specs,
            mrp=num(mrp), landing_price=num(mop), warranty="",
            source_file=F, source_sheet="NPC Specs",
        ))
    return out

records += parse_npc()

if __name__ == "__main__":
    by_sheet = {}
    for r in records:
        by_sheet[r["source_sheet"]] = by_sheet.get(r["source_sheet"], 0) + 1
    for s, c in by_sheet.items():
        print(s, c)
    print("TOTAL", len(records))
    with open(os.path.join(OUT, "samsung_full_records.json"), "w") as fh:
        json.dump(records, fh, indent=1, ensure_ascii=False)
