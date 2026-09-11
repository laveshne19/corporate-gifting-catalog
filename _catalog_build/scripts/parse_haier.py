# -*- coding: utf-8 -*-
import openpyxl, os, json

BASE = "/Users/laveshbansal/Downloads/📁 Master Folder/master price list"
OUT = os.path.join(BASE, "_catalog_build", "extracted")
F = "haier price list.xlsx"

CAT_MAP = {
    "AC": ("Home Appliances", "Air Conditioner"),
    "CE": ("Home Appliances", "Television"),
    "DF": ("Home Appliances", "Deep Freezer"),
    "MWO": ("Home Appliances", "Microwave Oven"),
    "REF": ("Home Appliances", "Refrigerator"),
    "WM": ("Home Appliances", "Washing Machine"),
    "WH": ("Home Appliances", "Water Heater"),
    "KA": ("Home Appliances", "Kitchen Chimney"),
    "RVC": ("Home Appliances", "Robotic Vacuum Cleaner"),
}

def num(v):
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return round(float(v), 2)
    return None

def clean(v):
    return str(v).strip() if v is not None else ""

def parse():
    wb = openpyxl.load_workbook(os.path.join(BASE, F), read_only=True, data_only=True)
    out = []
    for sheet in wb.sheetnames:
        cat, subcat = CAT_MAP.get(sheet, ("Home Appliances", sheet))
        ws = wb[sheet]
        rows = list(ws.iter_rows(values_only=True))
        header = [clean(h) for h in rows[1]] if len(rows) > 1 else []
        idx = {h: i for i, h in enumerate(header)}
        for row in rows[2:]:
            model = row[idx.get('Model', 0)] if idx.get('Model', 0) < len(row) else None
            if not model:
                continue
            mrp = num(row[idx['MRP']]) if 'MRP' in idx and idx['MRP'] < len(row) else None
            dp = num(row[idx['DP']]) if 'DP' in idx and idx['DP'] < len(row) else None
            out.append(dict(
                brand="Haier", category=cat, sub_category=subcat,
                product_name=f"Haier {clean(model)}", model_code=clean(model),
                description=f"Haier {subcat} — {clean(model)}", variant="", key_specs=subcat,
                mrp=mrp, landing_price=dp, warranty="",
                source_file=F, source_sheet=sheet,
            ))
    return out

if __name__ == "__main__":
    recs = parse()
    by_sheet = {}
    for r in recs:
        by_sheet[r["source_sheet"]] = by_sheet.get(r["source_sheet"], 0) + 1
    for s, c in by_sheet.items():
        print(s, c)
    print("TOTAL", len(recs))
    with open(os.path.join(OUT, "haier_records.json"), "w") as fh:
        json.dump(recs, fh, indent=1, ensure_ascii=False)
