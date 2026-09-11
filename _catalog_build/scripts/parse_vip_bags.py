# -*- coding: utf-8 -*-
import os, json, re, hashlib
from pptx import Presentation
from pptx.util import Emu

BASE = "/Users/laveshbansal/Downloads/📁 Master Folder/master price list"
OUT = os.path.join(BASE, "_catalog_build", "extracted")
IMG_OUT = os.path.join(BASE, "_catalog_build", "images", "VIP_Bags_Diwali_2026_pptx")
F = "VIP Bags -B2B Product Assortment-Diwali 2026.pptx"

os.makedirs(IMG_OUT, exist_ok=True)

def num(v):
    if v is None:
        return None
    s = str(v).replace(",", "")
    m = re.search(r'\d+(?:\.\d+)?', s)
    if not m:
        return None
    try:
        return round(float(m.group(0)), 2)
    except ValueError:
        return None

def clean(v):
    return re.sub(r'\s+', ' ', str(v)).strip()

p = Presentation(os.path.join(BASE, F))
records = []
cur_section = ""
seen_hash = set()

for si, slide in enumerate(p.slides):
    title = ""
    spec_text = ""
    table_rows = []
    images = []
    for shape in slide.shapes:
        if shape.has_text_frame:
            t = shape.text_frame.text.strip()
            if t:
                if "Specification" in t:
                    spec_text = t
                elif not title and len(t) < 80 and "Specification" not in t:
                    title = t
        if shape.has_table:
            tbl = shape.table
            rows = [[c.text.strip() for c in r.cells] for r in tbl.rows]
            table_rows.extend(rows)
        if shape.shape_type == 13:  # PICTURE
            try:
                img = shape.image
                w_px = int(Emu(shape.width).inches * 96)
                h_px = int(Emu(shape.height).inches * 96)
                images.append((img, w_px, h_px))
            except Exception:
                pass

    # section header slides (single short text, no table)
    if title and not table_rows and len(title) < 40:
        cur_section = title
        continue

    if not table_rows:
        continue

    header = [h.lower() for h in table_rows[0]]
    def col(name_opts):
        for i, h in enumerate(header):
            if any(h.startswith(n) for n in name_opts):
                return i
        return None
    idx_name = col(["product name"])
    idx_size = col(["size"])
    idx_brand = col(["brand"])
    idx_mrp = col(["mrp"])
    idx_rate = col(["net rate", "base rate"])

    if idx_name is None or idx_mrp is None:
        continue

    # save the largest embedded image on this slide once
    image_rel = ""
    if images:
        img, w, h = max(images, key=lambda t: t[1] * t[2])
        try:
            blob = img.blob
            hh = hashlib.md5(blob).hexdigest()[:12]
            if hh not in seen_hash:
                seen_hash.add(hh)
                ext = img.ext
                fname = f"slide{si:03d}_{hh}.{ext}"
                with open(os.path.join(IMG_OUT, fname), "wb") as fh:
                    fh.write(blob)
                image_rel = os.path.join("VIP_Bags_Diwali_2026_pptx", fname)
        except Exception:
            pass

    for row in table_rows[1:]:
        name = row[idx_name] if idx_name is not None else ""
        if not name:
            continue
        brand = row[idx_brand] if idx_brand is not None else "VIP"
        brand = "Skybags" if "sky" in brand.lower() else "VIP"
        size = row[idx_size] if idx_size is not None else ""
        mrp = num(row[idx_mrp])
        rate = num(row[idx_rate]) if idx_rate is not None else None
        full_name = f"{clean(name)}" + (f" ({size})" if size else "")
        records.append(dict(
            brand=brand, category="Luggage & Bags", sub_category=cur_section or title,
            product_name=full_name, model_code="", description=f"{clean(name)} — {clean(spec_text)}"[:400],
            variant=size, key_specs=clean(spec_text)[:300],
            mrp=mrp, landing_price=rate, warranty="",
            source_file=F, source_sheet=f"slide {si+1}",
            image_file=image_rel, image_source=f"Extracted from VIP Industries Diwali 2026 gifting deck (slide {si+1})" if image_rel else "",
        ))

if __name__ == "__main__":
    by_brand = {}
    for r in records:
        by_brand[r["brand"]] = by_brand.get(r["brand"], 0) + 1
    print(by_brand, "TOTAL", len(records))
    with_img = sum(1 for r in records if r["image_file"])
    print("with image:", with_img)
    with open(os.path.join(OUT, "vip_bags_records.json"), "w") as fh:
        json.dump(records, fh, indent=1, ensure_ascii=False)
    for r in records[:5]:
        print(r["product_name"], r["mrp"], r["landing_price"], r["image_file"])
