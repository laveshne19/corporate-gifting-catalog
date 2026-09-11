import openpyxl, glob, os, json, sys

BASE = "/Users/laveshbansal/Downloads/📁 Master Folder/master price list"
OUT = os.path.join(BASE, "_catalog_build", "extracted")

def dump_preview(f, max_rows=12):
    wb = openpyxl.load_workbook(f, read_only=True, data_only=True)
    result = {}
    for sheet in wb.sheetnames:
        ws = wb[sheet]
        rows = []
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i >= max_rows:
                break
            rows.append(list(row))
        result[sheet] = rows
    return result

if __name__ == "__main__":
    files = sorted(glob.glob(os.path.join(BASE, "*.xlsx")))
    report_lines = []
    for f in files:
        name = os.path.basename(f)
        report_lines.append(f"\n{'='*100}\nFILE: {name}\n{'='*100}")
        try:
            preview = dump_preview(f)
            for sheet, rows in preview.items():
                report_lines.append(f"\n--- Sheet: {sheet} ---")
                for r in rows:
                    report_lines.append(str(r))
        except Exception as e:
            report_lines.append(f"ERROR: {e}")
    out_path = os.path.join(OUT, "xlsx_preview.txt")
    with open(out_path, "w") as fh:
        fh.write("\n".join(report_lines))
    print("Wrote", out_path)
