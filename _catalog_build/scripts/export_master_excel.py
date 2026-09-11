# -*- coding: utf-8 -*-
"""
Exports the full master catalog to a single sortable/filterable Excel
workbook — brand, product name, category, MRP, landing cost (NLC) and the
website's displayed budget range side by side, so the whole catalog can be
reviewed and sorted outside the site.
"""
import json, os, re
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

ILLEGAL_CHARS_RE = re.compile(r'[\x00-\x08\x0b\x0c\x0e-\x1f]')
def clean(v):
    if isinstance(v, str):
        return ILLEGAL_CHARS_RE.sub('', v)
    return v

BASE = "/Users/laveshbansal/Downloads/📁 Master Folder/master price list"
OUT_DIR = os.path.join(BASE, "_catalog_build", "output")
SRC = os.path.join(OUT_DIR, "master_consolidated.json")
DEST = os.path.join(BASE, "Master_Product_List.xlsx")

COLUMNS = [
    ("Product ID", "product_id", 12),
    ("Brand", "brand", 18),
    ("Product Name", "product_name", 42),
    ("Category", "category", 22),
    ("Sub-Category", "sub_category", 20),
    ("Variant", "variant", 16),
    ("Model Code", "model_code", 16),
    ("MRP (Rs.)", "mrp", 12),
    ("Landing Cost / NLC (Rs.)", "landing_price", 16),
    ("Website Budget Low (Rs.)", "budget_low", 14),
    ("Website Budget High (Rs.)", "budget_high", 14),
    ("Budget Basis", "budget_basis", 20),
    ("Has Image", "_has_image", 10),
    ("Source Price List", "source_file", 34),
]

def main():
    recs = json.load(open(SRC))
    recs = sorted(recs, key=lambda r: (r.get("brand") or "", r.get("product_name") or ""))

    wb = Workbook()
    ws = wb.active
    ws.title = "Master Product List"

    header_fill = PatternFill(start_color="0F2A4A", end_color="0F2A4A", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True, size=11)
    for i, (label, _, width) in enumerate(COLUMNS, 1):
        cell = ws.cell(row=1, column=i, value=label)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(vertical="center")
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = "A2"

    for r_i, r in enumerate(recs, 2):
        r["_has_image"] = "Yes" if r.get("image_file") else "No"
        for c_i, (_, key, _w) in enumerate(COLUMNS, 1):
            val = r.get(key)
            if key in ("mrp", "landing_price", "budget_low", "budget_high") and val is not None:
                val = round(float(val), 2)
            val = clean(val)
            ws.cell(row=r_i, column=c_i, value=val if val not in (None, "") else None)

    last_row = len(recs) + 1
    last_col = get_column_letter(len(COLUMNS))
    ws.auto_filter.ref = f"A1:{last_col}{last_row}"

    # Summary sheet
    ws2 = wb.create_sheet("Summary")
    from collections import Counter
    brand_counts = Counter(r["brand"] for r in recs)
    cat_counts = Counter(r["category"] for r in recs)
    missing_img = sum(1 for r in recs if not r.get("image_file"))
    ws2.append(["Master Catalog Summary"])
    ws2["A1"].font = Font(bold=True, size=14)
    ws2.append([])
    ws2.append(["Total products", len(recs)])
    ws2.append(["Total brands", len(brand_counts)])
    ws2.append(["Total categories", len(cat_counts)])
    ws2.append(["Products missing an image", missing_img])
    ws2.append([])
    ws2.append(["Brand", "Product Count"])
    for cell in ws2[7]:
        cell.font = Font(bold=True)
    for b, n in brand_counts.most_common():
        ws2.append([b, n])
    ws2.column_dimensions["A"].width = 28
    ws2.column_dimensions["B"].width = 16

    wb.save(DEST)
    print(f"Wrote {DEST} — {len(recs)} products, {len(COLUMNS)} columns")

if __name__ == "__main__":
    main()
